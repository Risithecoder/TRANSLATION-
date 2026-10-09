from docx import Document
from docx.oxml import parse_xml
import xml.etree.ElementTree as ET

# 1. Create a docx with math for testing
doc = Document()
p = doc.add_paragraph('Before math. ')
math_xml = """<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>a=b</m:t></m:r></m:oMath>"""
parsed_math = parse_xml(math_xml)
p._element.append(parsed_math)
p.add_run(' After math.')
doc.save('test_math_2.docx')

# 2. Extract math and replace with placeholder
doc2 = Document('test_math_2.docx')
math_elements = []
for p in doc2.paragraphs:
    # Find all m:oMath elements in the paragraph
    # The namespace for math is http://schemas.openxmlformats.org/officeDocument/2006/math
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    for math_el in p._element.findall('.//m:oMath', ns):
        math_elements.append(math_el)
        # Create a text run placeholder
        placeholder = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:t>[MATH_{len(math_elements)-1}]</w:t></w:r>')
        # Replace the math_el with placeholder
        parent = math_el.getparent()
        parent.replace(math_el, placeholder)

    # Also handle block level math m:oMathPara
    
    for math_para in p._element.findall('.//m:oMathPara', ns):
        math_elements.append(math_para)
        placeholder = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:t>[MATH_{len(math_elements)-1}]</w:t></w:r>')
        parent = math_para.getparent()
        parent.replace(math_para, placeholder)

doc2.save('test_math_2_replaced.docx')

import mammoth
with open('test_math_2_replaced.docx', 'rb') as f:
    result = mammoth.convert_to_html(f)
    print("HTML with placeholders:", result.value)

# 3. Re-insert math into a new docx
doc3 = Document()
p3 = doc3.add_paragraph('Re-inserted: ')
p3._element.append(math_elements[0])
doc3.save('test_math_2_restored.docx')
print("Successfully restored to a new document.")
